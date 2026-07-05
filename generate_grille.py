"""
Génère une grille horaire visuelle (un onglet par jour) à partir :
  - de data/horaires_paleo.xlsx     : les horaires des concerts, par scène
  - de output/planning_<Jour>.xlsx  : les plannings d'équipe déjà générés
    par generate_planning.py (un par jour)

Pour chaque jour disponible (présent à la fois dans horaires_paleo.xlsx et
sous forme de output/planning_<Jour>.xlsx déjà généré), produit un onglet
avec :
  - une ligne de bandeau horaire (de 15h à 03h, par pas de 5 minutes) ;
  - une ligne par scène, avec le nom de chaque concert positionné et
    dimensionné selon son horaire réel ;
  - une ligne par tâche du planning, avec les personnes affectées
    positionnées et dimensionnées selon l'horaire de chaque créneau.

Les jours dont le planning n'a pas encore été généré (output/planning_<Jour>.xlsx
manquant) sont ignorés, avec un avertissement.

Utilisation :
    python generate_grille.py
    -> écrit output/Grille_horaire.xlsx
"""

import datetime as dt
from collections import defaultdict

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins

from generate_planning import DATA_DIR, JOURS_VALIDES, OUTPUT_DIR, to_minutes

JOURS_FR = ["Lundi", "Mardi", "Mercredi", "Jeudi", "Vendredi", "Samedi", "Dimanche"]
MOIS_FR = [
    "janvier", "février", "mars", "avril", "mai", "juin",
    "juillet", "août", "septembre", "octobre", "novembre", "décembre",
]

# La grille couvre de 15h à 03h le lendemain, par pas de 5 minutes. Toute
# heure antérieure à midi est considérée comme appartenant au petit matin
# du jour suivant (continuité de la grille après minuit).
PIVOT_NOCTURNE_MINUTES = 12 * 60
GRID_START_MINUTES = 15 * 60
GRID_END_MINUTES = (24 + 3) * 60
MINUTES_PAR_COLONNE = 5
COL_DEBUT_GRILLE = 3
COL_FIN_GRILLE = COL_DEBUT_GRILLE + (GRID_END_MINUTES - GRID_START_MINUTES) // MINUTES_PAR_COLONNE - 1

# Palette noir & blanc (économise l'encre à l'impression) : pas d'aplats
# sombres, juste des nuances de gris clair pour distinguer les zones, et du
# texte noir partout.
FOND_LABEL = PatternFill("solid", fgColor="D9D9D9")
FOND_ENTETE = PatternFill("solid", fgColor="BFBFBF")
FOND_TIMELINE = PatternFill("solid", fgColor="FFFFFF")
FOND_CONCERT = PatternFill("solid", fgColor="D9D9D9")
FOND_PLANNING = PatternFill("solid", fgColor="F2F2F2")
TEXTE_LABEL = Font(color="000000", bold=True)
TEXTE_HEURE = Font(color="000000", size=8)
TEXTE_CONCERT = Font(color="000000", bold=True, size=8)
TEXTE_PLANNING = Font(color="000000", size=7)

# Quadrillage de la timeline : trait fin entre les colonnes de 5 minutes,
# trait plus marqué à chaque heure pleine, pour pouvoir lire les horaires
# avec exactitude.
BORDURE_FINE = Side(style="thin", color="BFBFBF")
BORDURE_HEURE = Side(style="thin", color="000000")


def bordure_pour_colonne(col: int) -> Border:
    cote_gauche = BORDURE_HEURE if (col - COL_DEBUT_GRILLE) % 12 == 0 else BORDURE_FINE
    return Border(left=cote_gauche, right=BORDURE_FINE, top=BORDURE_FINE, bottom=BORDURE_FINE)


def minutes_ajustees(valeur_brute) -> int:
    """Heure brute -> minutes depuis minuit, repoussée de 24h si elle
    appartient en réalité au petit matin du jour suivant."""
    minutes = to_minutes(valeur_brute)
    return minutes + 24 * 60 if minutes < PIVOT_NOCTURNE_MINUTES else minutes


def plage_minutes(debut_brut, fin_brut) -> tuple:
    debut = minutes_ajustees(debut_brut)
    fin = minutes_ajustees(fin_brut)
    if fin <= debut:
        fin += 24 * 60
    return debut, fin


def colonne_pour_minutes(minutes: int) -> int:
    return COL_DEBUT_GRILLE + (minutes - GRID_START_MINUTES) // MINUTES_PAR_COLONNE


def plage_colonnes(debut_minutes: int, fin_minutes: int):
    """Renvoie (col_debut, col_fin) bornées à la grille, ou None si la
    plage tombe entièrement hors de la grille."""
    col_debut = max(COL_DEBUT_GRILLE, colonne_pour_minutes(debut_minutes))
    col_fin = min(COL_FIN_GRILLE, colonne_pour_minutes(fin_minutes) - 1)
    if col_fin < col_debut:
        return None
    return col_debut, col_fin


# ---------------------------------------------------------------------------
# Chargement des horaires de concerts (data/horaires_paleo.xlsx)
# ---------------------------------------------------------------------------

def charger_horaires_concerts(path):
    """Lit le tableau empilé (une scène, puis une date, puis ses concerts,
    répété pour chaque scène) et renvoie (scenes, concerts) où :
      - scenes : liste ordonnée des noms de scène ;
      - concerts : dict (scène, date) -> liste de (artiste, début_min, fin_min).
    """
    df = pd.read_excel(path, sheet_name="Sheet1", header=None)

    scenes = []
    concerts = defaultdict(list)
    scene_courante = None
    date_courante = None

    for _, row in df.iterrows():
        col0, col1, col2 = row[0], row[1], row[2]

        if pd.isna(col0):
            scene_courante = None
            date_courante = None
            continue

        if isinstance(col0, (dt.datetime, pd.Timestamp)):
            date_courante = col0.date()
            continue

        if pd.isna(col1) and pd.isna(col2):
            if str(col0).strip().lower() != "band":
                scene_courante = str(col0).strip()
                if scene_courante not in scenes:
                    scenes.append(scene_courante)
            continue

        if scene_courante is None or date_courante is None:
            continue
        if pd.isna(col1) or pd.isna(col2):
            continue

        debut, fin = plage_minutes(col1, col2)
        concerts[(scene_courante, date_courante)].append((str(col0).strip(), debut, fin))

    return scenes, concerts


def construire_dates_par_jour(concerts):
    """Déduit, à partir des dates trouvées dans horaires_paleo.xlsx, la date
    calendaire de chaque jour de la semaine (Mardi -> 2026-07-21, etc.)."""
    dates = sorted({date for (_, date) in concerts.keys()})
    return {JOURS_FR[date.weekday()]: date for date in dates}


# ---------------------------------------------------------------------------
# Chargement des tâches déjà planifiées (output/planning_<Jour>.xlsx)
# ---------------------------------------------------------------------------

def charger_taches_planning(planning_path, jour):
    """Renvoie un dict tache -> liste de (personnes, début_min, fin_min),
    dans l'ordre d'apparition des tâches dans le planning."""
    wb = load_workbook(planning_path, data_only=True)
    ws = wb[jour]

    taches = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None or str(row[0]).startswith("("):
            continue
        tache, debut, fin, _requis, _affectes, _manquant, personnes = row[:7]
        taches.setdefault(str(tache).strip(), [])
        if not personnes:
            continue
        debut_min, fin_min = plage_minutes(debut, fin)
        taches[str(tache).strip()].append((str(personnes).strip(), debut_min, fin_min))

    return taches


# ---------------------------------------------------------------------------
# Construction de la grille
# ---------------------------------------------------------------------------

def ecrire_entete_heures(ws, ligne):
    ws.cell(row=ligne, column=1).fill = FOND_ENTETE
    minutes = GRID_START_MINUTES
    while minutes < GRID_END_MINUTES:
        col = colonne_pour_minutes(minutes)
        cell = ws.cell(row=ligne, column=col, value=f"{(minutes % (24 * 60)) // 60:02d}h")
        cell.font = TEXTE_HEURE
        minutes += 60
    for col in range(COL_DEBUT_GRILLE, COL_FIN_GRILLE + 1):
        cell = ws.cell(row=ligne, column=col)
        cell.fill = FOND_ENTETE
        cell.border = bordure_pour_colonne(col)
    ws.row_dimensions[ligne].height = 14


def ecrire_ligne(ws, ligne, label, evenements, fond_evenement, texte_evenement, hauteur=20):
    cell_label = ws.cell(row=ligne, column=1, value=label)
    cell_label.fill = FOND_LABEL
    cell_label.font = TEXTE_LABEL
    cell_label.alignment = Alignment(horizontal="right", vertical="center")

    for col in range(COL_DEBUT_GRILLE, COL_FIN_GRILLE + 1):
        cell = ws.cell(row=ligne, column=col)
        cell.fill = FOND_TIMELINE
        cell.border = bordure_pour_colonne(col)

    for texte, debut_min, fin_min in evenements:
        plage = plage_colonnes(debut_min, fin_min)
        if plage is None:
            continue
        col_debut, col_fin = plage
        if col_fin > col_debut:
            ws.merge_cells(start_row=ligne, start_column=col_debut, end_row=ligne, end_column=col_fin)
        for col in range(col_debut, col_fin + 1):
            cell = ws.cell(row=ligne, column=col)
            cell.fill = fond_evenement
            cell.border = bordure_pour_colonne(col)
        cell_ancre = ws.cell(row=ligne, column=col_debut, value=texte)
        cell_ancre.font = texte_evenement
        cell_ancre.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.row_dimensions[ligne].height = hauteur


def assigner_lanes(evenements):
    """Répartit des événements qui peuvent se chevaucher (ex: une animation
    continue qui couvre plusieurs concerts ponctuels) sur autant de lignes
    ("lanes") que nécessaire, afin de ne jamais fusionner deux cellules qui
    se chevauchent. Renvoie une liste de lignes, chacune une liste
    d'événements ; toujours au moins une ligne (éventuellement vide)."""
    lanes = []
    fins_lanes = []
    for evt in sorted(evenements, key=lambda e: e[1]):
        _, debut, _ = evt
        # On regarde les lignes existantes de la plus récemment utilisée à
        # la plus ancienne : les évènements courts et consécutifs
        # continuent ainsi de remplir la même ligne, plutôt que de
        # reprendre une ligne juste libérée par un évènement isolé très
        # long (ex: une animation continue).
        for i in range(len(fins_lanes) - 1, -1, -1):
            if fins_lanes[i] <= debut:
                lanes[i].append(evt)
                fins_lanes[i] = evt[2]
                break
        else:
            lanes.append([evt])
            fins_lanes.append(evt[2])
    lanes = lanes or [[]]
    # Les lignes les plus "denses" (le plus d'événements, donc les plus
    # courts) passent en premier ; une ligne avec peu d'événements longs
    # (ex: une animation continue) passe en dessous.
    lanes.sort(key=len, reverse=True)
    return lanes


def hauteur_pour_lanes(hauteur_standard, nb_lanes: int):
    if nb_lanes <= 1:
        return hauteur_standard
    return hauteur_standard * 2 / 3


def ajuster_largeurs_colonnes(ws):
    ws.column_dimensions["A"].width = 22
    for col in range(2, COL_FIN_GRILLE + 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = 1.5


def configurer_impression(ws):
    """Met en page pour une impression sur papier A4 : paysage, marges
    réduites, mise à l'échelle pour tenir sur la hauteur d'une page (la
    largeur s'étale sur autant de pages que nécessaire, à la même échelle),
    et la colonne des libellés répétée sur chaque page imprimée."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 0
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.2, right=0.2, top=0.3, bottom=0.3, header=0, footer=0)
    ws.print_title_cols = "A:A"
    ws.print_options.horizontalCentered = True


# Disposition personnalisée des lignes, du haut vers le bas, regroupées en
# sections séparées par un bandeau d'heures répété (lisibilité sur une
# feuille haute). "tache_prefixe" développe toutes les tâches dont le nom
# commence par ce préfixe, dans leur ordre d'apparition dans le planning.
DISPOSITION = [
    [("scene", "La Ruche"), ("tache", "Abeilles Ruche"), ("tache", "Entrée backstage")],
    [("tache", "Ailes"), ("tache", "Bar public"), ("tache", "Bar loges")],
    [("scene", "Grande Scène"), ("scene", "Véga"), ("scene", "Belleville"), ("scene", "Dôme"), ("scene", "Club Tent")],
    [("tache_prefixe", "tech"), ("tache_prefixe", "déambule")],
]


def resoudre_section(section, scenes, taches, deja_places):
    """Développe une section de DISPOSITION en une liste de (type, nom) à
    dessiner, en ignorant ce qui ne correspond à rien et en notant ce qui a
    été placé (pour le rattrapage des scènes/tâches non listées)."""
    elements = []
    for type_, valeur in section:
        if type_ == "scene":
            noms = [s for s in scenes if s.lower() == valeur.lower()]
        elif type_ == "tache":
            noms = [t for t in taches if t.lower() == valeur.lower()]
        else:  # tache_prefixe
            type_ = "tache"
            noms = [t for t in taches if t.lower().startswith(valeur.lower())]
        for nom in noms:
            elements.append((type_, nom))
            deja_places.add((type_, nom.lower()))
    return elements


def construire_grille(output_path):
    scenes, concerts = charger_horaires_concerts(DATA_DIR / "horaires_paleo.xlsx")
    dates_par_jour = construire_dates_par_jour(concerts)

    wb = Workbook()
    wb.remove(wb.active)

    for jour in JOURS_FR:
        if jour.lower() not in JOURS_VALIDES or jour not in dates_par_jour:
            continue

        planning_path = OUTPUT_DIR / f"planning_{jour}.xlsx"
        if not planning_path.exists():
            print(f"Attention : {planning_path.name} introuvable, jour {jour} ignoré.")
            continue

        date = dates_par_jour[jour]
        titre = f"{jour} {date.day} {MOIS_FR[date.month - 1]} {date.year}"
        ws = wb.create_sheet(title=titre[:31])

        titre_cell = ws.cell(row=1, column=1, value=titre)
        titre_cell.font = Font(bold=True, size=14, color="000000")
        titre_cell.fill = FOND_LABEL
        ws.row_dimensions[1].height = 20

        taches = charger_taches_planning(planning_path, jour)
        deja_places = set()

        ligne = 2
        ecrire_entete_heures(ws, ligne)
        ligne += 1
        for section in DISPOSITION:
            for type_, nom in resoudre_section(section, scenes, taches, deja_places):
                if type_ == "scene":
                    lanes = assigner_lanes(concerts.get((nom, date), []))
                    hauteur = hauteur_pour_lanes(26, len(lanes))
                    fond, texte = FOND_CONCERT, TEXTE_CONCERT
                else:
                    lanes = assigner_lanes(taches.get(nom, []))
                    hauteur = hauteur_pour_lanes(20, len(lanes))
                    fond, texte = FOND_PLANNING, TEXTE_PLANNING
                for i, lane in enumerate(lanes):
                    ecrire_ligne(ws, ligne, nom if i == 0 else "", lane, fond, texte, hauteur=hauteur)
                    ligne += 1
            ecrire_entete_heures(ws, ligne)
            ligne += 1

        restantes = [
            (type_, nom) for type_, nom in
            ([("scene", s) for s in scenes] + [("tache", t) for t in taches])
            if (type_, nom.lower()) not in deja_places
        ]
        for type_, nom in restantes:
            if type_ == "scene":
                lanes = assigner_lanes(concerts.get((nom, date), []))
                hauteur = hauteur_pour_lanes(26, len(lanes))
                fond, texte = FOND_CONCERT, TEXTE_CONCERT
            else:
                lanes = assigner_lanes(taches.get(nom, []))
                hauteur = hauteur_pour_lanes(20, len(lanes))
                fond, texte = FOND_PLANNING, TEXTE_PLANNING
            for i, lane in enumerate(lanes):
                ecrire_ligne(ws, ligne, nom if i == 0 else "", lane, fond, texte, hauteur=hauteur)
                ligne += 1
        if restantes:
            ecrire_entete_heures(ws, ligne)
            ligne += 1

        ajuster_largeurs_colonnes(ws)
        configurer_impression(ws)
        ws.freeze_panes = "C3"

    if not wb.sheetnames:
        raise RuntimeError(
            "Aucun jour n'a pu être construit : vérifiez que des fichiers "
            "output/planning_<Jour>.xlsx existent pour au moins un jour présent "
            "dans data/horaires_paleo.xlsx."
        )

    OUTPUT_DIR.mkdir(exist_ok=True)
    wb.save(output_path)
    return output_path


def main():
    output_path = OUTPUT_DIR / "Grille_horaire.xlsx"
    construire_grille(output_path)
    print(f"Grille horaire générée : {output_path}")


if __name__ == "__main__":
    main()
